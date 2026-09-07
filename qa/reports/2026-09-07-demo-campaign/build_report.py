from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

base = r"C:\Users\USER\Desktop\CodentTech\cleercut\frontend\qa\reports\2026-09-07-demo-campaign"
shots = os.path.join(base, "screenshots")
out = os.path.join(base, "CleerCut-Demo-Campaign-Test-Report.docx")

doc = Document()


def shade_paragraph(paragraph, hex_color):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)


def add_status(text, passed=True):
    p = doc.add_paragraph()
    run = p.add_run("PASSED" if passed else "FAILED")
    run.bold = True
    run.font.color.rgb = (
        RGBColor(0x16, 0x65, 0x34) if passed else RGBColor(0x99, 0x1B, 0x1B)
    )
    p.add_run(f"  —  {text}")
    shade_paragraph(p, "E8F5E9" if passed else "FFEBEE")
    return p


doc.add_heading("CleerCut — Fixed Demo Campaign Test Report", level=0)
doc.add_paragraph("Date: 7 September 2026")
doc.add_paragraph("Environment: Local (localhost:3000 / localhost:5000)")
doc.add_paragraph(
    "Account tested: theasadabbasdev+3test3@gmail.com (Gabriel / brand)"
)
doc.add_paragraph(
    "Purpose: Verify the Fixed Demo Campaign experience against the product spec — populated Applications / Active / Completed, Sales & ROI report, mutation blocks, isolation, and archive behavior."
)

overall = doc.add_paragraph()
overall.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = overall.add_run("ALL CHECKS PASSED")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = RGBColor(0x16, 0x65, 0x34)
shade_paragraph(overall, "C8E6C9")

doc.add_heading("What We Checked", level=1)
doc.add_paragraph(
    "We walked the brand campaign hub with the Demo Campaign selected: Applications with fabricated applicants and messages, Active hired creators, Completed creators, and the full completed report including Sales & ROI. "
    "We attempted hire send, reject confirm, direct message send, calendar scheduling, and PDF export to confirm those writes stay blocked with clear on-screen messages. "
    "We confirmed demo creators do not appear in Discover+, and that archiving the demo removes it from the campaign list when the brand already has three or more real campaigns."
)

doc.add_heading("Page / Feature Test Results", level=1)
add_status(
    "Demo campaign available in Multi-Creator dropdown as “Demo Campaign — Verified Creators Sample”"
)
add_status(
    "Persistent banner shown: “Demo Campaign — creators and data shown are fictional.”"
)
add_status("Applications tab populated with fictional applicants and application messages")
add_status("Active tab populated with hired demo creators")
add_status("Completed tab populated with completed demo creators")
add_status("Completed report opens and Overview shows believable demo metrics")
add_status(
    "Sales & ROI tab shows revenue, orders, ROI, and sample Shopify-style sales figures (no live store required)"
)
add_status("Demo banner visible inside the report viewer")
add_status("Reject confirm blocked with inline demo message; applicant remains")
add_status("Send Offer blocked with “Sending is disabled for the sample campaign.”")
add_status("Message compose works; Send blocked with demo messaging")
add_status(
    "Calendar opens for browsing; scheduling shown as disabled for the sample campaign"
)
add_status("Export as PDF blocked with “PDF export is disabled for the sample campaign.”")
add_status("Demo creator social links are inert (no Instagram/TikTok navigation)")
add_status(
    "Discover+ search for demo creator returns no creators (demo creators stay brand-only)"
)
add_status(
    "Demo campaign archives when brand already has 3+ real campaigns; removed from dropdown"
)

doc.add_heading("Supporting Files / Related Checks", level=1)
add_status(
    "Database migration for is_demo flags applied successfully before functional checks"
)
add_status("Backend restarted with Demo Campaign module loaded")
p = doc.add_paragraph()
run = p.add_run("NOTE")
run.bold = True
p.add_run(
    "  —  This brand already had 22 real campaigns, so the product correctly would not auto-seed a demo. "
    "For this QA pass the demo was seeded once for verification, then archived with the same archive rule used when a brand reaches three real campaigns. "
    "Auto-select on a brand with zero real campaigns was not re-tested on this account."
)

doc.add_heading("Key Confirmation", level=1)
doc.add_paragraph(
    "Sales & ROI is viewable on the demo completed report with fictional but realistic numbers, while mutating actions (send offer, reject, message send, calendar write, PDF export) stay blocked at submit time."
)

doc.add_heading("What This Means", level=1)
doc.add_paragraph(
    "New and early-stage brands can explore a full CleerCut campaign funnel — including the ROI story — without empty tabs and without any of those sample actions changing real inbox, payments, Shopify, or reporting aggregates."
)

doc.add_heading("Sample Screenshots", level=1)
shots_meta = [
    ("step-02-applications-demo.png", "1. Applications — demo applicants and banner"),
    ("step-04-reject-confirm-blocked.png", "2. Reject blocked on confirm"),
    ("step-05-send-offer-blocked.png", "3. Send Offer blocked"),
    ("step-06-message-send-blocked.png", "4. Message send blocked"),
    ("step-07-active-demo.png", "5. Active — hired demo creators"),
    ("step-08-calendar-blocked.png", "6. Calendar scheduling disabled"),
    ("step-10-report-sales-roi.png", "7. Report — Sales & ROI"),
    ("step-12-report-banner-fixed.png", "8. Report — demo banner"),
    ("step-13-discover-no-demo.png", "9. Discover+ excludes demo creators"),
    ("step-14-applications-social-inert.png", "10. Social links inert on demo cards"),
    ("step-15-demo-archived.png", "11. Demo removed after archive"),
]
for fname, caption in shots_meta:
    path = os.path.join(shots, fname)
    if not os.path.exists(path):
        continue
    doc.add_heading(caption, level=2)
    doc.add_picture(path, width=Inches(6.2))
    cap = doc.add_paragraph(caption)
    if cap.runs:
        cap.runs[0].italic = True

doc.add_heading("Final Status for Client", level=1)
final = doc.add_paragraph()
final.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = final.add_run(
    "PASSED — Fixed Demo Campaign checks completed successfully on the tested brand account."
)
fr.bold = True
fr.font.color.rgb = RGBColor(0x16, 0x65, 0x34)
shade_paragraph(final, "C8E6C9")

doc.save(out)
print("Wrote", out)
